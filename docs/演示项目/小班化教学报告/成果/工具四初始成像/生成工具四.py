"""按专家原表固定字段填制工具四；HTML、Word、文本共用同一份内容。"""
from pathlib import Path
from html import escape
from zipfile import ZipFile
import base64, csv, json, re, xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent
ROOT = next(p for p in OUT.parents if (p / 'mobile-app').exists())
DATA = json.loads((OUT / '工具四内容.json').read_text())
S = json.loads((OUT / '工具四模板结构.json').read_text())
RULES = json.loads((ROOT / 'prompt/student/学生初始画像档案/工具四字段规则-v1.json').read_text())
FIVE_LABELS = [r['canonicalFieldLabel'] for r in RULES['rules'] if r['ruleKey'].startswith('card.five.')]
assert set(FIVE_LABELS) == set(S['five_labels'])
assert [label[0] for label in FIVE_LABELS] == RULES['fiveEducationPolicy']['displayOrder']
CSV = Path('/Users/mayday/Downloads/川大附小2026级_谌芷晞_李林卿_吴浩_问卷与成长档案.csv')
P = S['paragraphs']
TITLE = P[0]
H1 = next(p for p in P if p.startswith('一、'))
H2 = next(p for p in P if p.startswith('二、'))
FIVE = next(p for p in P if p.startswith('五育初印象'))
THREE = next(p for p in P if p.startswith('三力小萌芽'))
LIGHT = next(p for p in P if p.startswith('光芒图初始方向'))
WATCH = next(p for p in P if p.startswith('需要留意'))
WATCH_HELP = next(p for p in P if p.startswith('三方信息不一致'))
MESSAGE = next(p for p in P if p.startswith('成像寄语'))
LIGHT_LABELS = [next(p for p in P if p.startswith(x)).split('：')[0] + '：' for x in ['孩子自己最想要', '孩子觉得自己最厉害', '家长最希望', '教师观察到']]
LIGHT_RESULT = next(p for p in P if p.startswith('✨')).removeprefix('✨ ')
LIGHT_NOTE = next(p for p in P if p.startswith('（此为'))
MESSAGE_LABELS = [next(p for p in P if p.startswith(x)) for x in ['这个孩子', '他/她', '我对接下来']]
SOURCE_LABELS = ['家长说……', '孩子说……', '老师看到了……']
MATRIX_LABELS = [c.replace('\n','') for c in S['tables'][0][0]]


def validate():
    # 原表重新读取并逐字比对，防止模板结构在后续迭代中漂移。
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(S['source']) as z: root=ET.fromstring(z.read('word/document.xml'))
    def t(el): return ''.join(x.text or '' for x in el.findall('.//w:t',ns))
    paragraphs=[]; tables=[]
    for el in root.find('w:body',ns):
        if el.tag.endswith('}p'): paragraphs.append(t(el))
        elif el.tag.endswith('}tbl'):
            tables.append([['\n'.join(t(p) for p in c.findall('w:p',ns)) for c in row.findall('w:tc',ns)] for row in el.findall('w:tr',ns)])
    assert paragraphs==P and tables==S['tables'], '工具四原表已变化，请重新核对'
    rows=list(csv.DictReader(CSV.open(encoding='utf-8-sig',newline='')))
    for name,d in DATA.items():
        rr=[r for r in rows if r['学生']==name]
        assert {r['名称'] for r in rr}=={'家长问卷（家长眼中的孩子）','2026级师生对话访谈采集','2026级新生入学观察记录表（老师眼中的孩子）'}
        assert [m['dim'] for m in d['matrix']]==S['dimensions']
        assert d['matrix'][0]['insight'] in S['judgment_options']
        for m,prefix in zip(d['matrix'][1:],S['judgment_prefixes']):assert m['insight'].startswith(prefix)
        assert [x[0] for x in d['five']]==FIVE_LABELS
        assert [x[0] for x in d['three']]==S['three_labels']
        assert len(d['message']['traits'])==3
        source_names={'parent':'家长问卷（家长眼中的孩子）','child':'2026级师生对话访谈采集','teacher':'2026级新生入学观察记录表（老师眼中的孩子）'}
        for key,source in source_names.items():
            answers={r['回答'] for r in rr if r['名称']==source}
            assert 1<=len(d['quotes'][key])<=2
            for quote in d['quotes'][key]:assert (quote['quote'] if isinstance(quote,dict) else quote) in answers,(name,key,quote)
        txt=json.dumps(d,ensure_ascii=False)
        for other in DATA:
            if other!=name:assert other not in txt,(name,other)
    (OUT/'内容核对.json').write_text(json.dumps({'原表结构逐字核对':True,'学生数':len(DATA),'气质判断允许值':S['judgment_options'],'其他判断固定前缀':S['judgment_prefixes'],'原表五育名称':S['five_labels'],'五育名称':FIVE_LABELS,'五育修订依据':RULES['fiveEducationPolicy']['revision'],'三力名称':S['three_labels'],'三方引文已按来源核对':True,'使用日常评价记录':False},ensure_ascii=False,indent=2))


def e(v):return escape(str(v),quote=True)
def image_uri(d):
    path=ROOT/f'mobile-app/assets/resources/term-report-covers/primary-lower-{d["gender"]}-discovery-v2.png'
    return path,'data:image/png;base64,'+base64.b64encode(path.read_bytes()).decode()
def heading(text,tag='h2'):
    # 拆分只影响视觉层级，文字保持原表完整。
    i=text.find('（')
    main,note=(text[:i],text[i:]) if i>=0 else (text,'')
    return f'<header class="section-heading"><{tag}>{e(main)}</{tag}>'+ (f'<p class="template-note">{e(note)}</p>' if note else '')+'</header>'
def field(label,value,cls=''):
    return f'<div class="field {cls}"><dt>{e(label)}</dt><dd>{e(value)}</dd></div>'
def quote_text(q):return q['context']+'：“'+q['quote']+'”。' if isinstance(q,dict) else '“'+q+'”'

def message_texts(d):
    m=d['message']
    a=MESSAGE_LABELS[0]
    for trait in m['traits']:a=a.replace('________',trait,1)
    b=MESSAGE_LABELS[1].replace('________',m['need'])
    c=MESSAGE_LABELS[2].replace('________',m['begin'])
    return [a,b,c]

CSS = r'''
:root{--ink:#12365c;--body:#263e55;--muted:#50677c;--line:#dce6f0;--paper:#fff;--blue:#eef5fc;--green:#edf8f4;--green-ink:#17644e;--violet:#f2f2fc;--s1:4px;--s2:8px;--s3:12px;--s4:16px;--s5:24px;--s6:32px}
*{box-sizing:border-box}body{margin:0;background:#eaf0f7;color:var(--body);font-family:"PingFang SC","Microsoft YaHei",sans-serif;font-size:14px;font-weight:400;line-height:1.65}h1,h2,h3,p,dl,dd,blockquote{margin:0}b,strong,h1,h2,h3{font-weight:700}button{font:inherit;font-weight:700;min-height:48px;padding:12px 20px;color:#fff;background:var(--ink);border:0;border-radius:12px;cursor:pointer}button:hover{background:#215783}button:focus-visible{outline:3px solid #215783;outline-offset:4px}.toolbar{position:sticky;top:0;z-index:10;display:flex;align-items:center;justify-content:space-between;gap:16px;padding:12px 24px;background:#fff;border-bottom:1px solid var(--line)}.toolbar strong{color:var(--ink)}
.page{position:relative;width:210mm;height:297mm;margin:24px auto;padding:16mm 16mm 18mm;background:#fff;box-shadow:0 4px 16px #d8e2ec}.page footer{position:absolute;bottom:7mm;right:16mm;font-size:12px;line-height:1;color:var(--muted)}h2{font-size:28px;line-height:1.3;color:var(--ink)}h3{font-size:20px;line-height:1.4;color:var(--ink)}.section-heading{margin-bottom:24px}.template-note{font-size:12px;line-height:1.65;color:var(--muted);margin-top:8px}.section-heading.sub{margin-top:32px}.field dt{font-size:14px;font-weight:700;color:var(--muted);margin-bottom:8px}.field dd{font-size:16px;line-height:1.65;color:var(--ink)}
.cover{display:flex;flex-direction:column;background:linear-gradient(150deg,#fff 55%,#f7fbff 88%,#fff)}.cover.girl{background:linear-gradient(150deg,#fff 55%,#fbf9ff 88%,#fff)}.school{font-size:20px;color:var(--muted);font-weight:700}.cover h1{font-size:52px;line-height:1.3;color:var(--ink);margin-top:40px}.cover h1 span{display:block}.cover .sources{font-size:16px;margin-top:16px;color:var(--muted)}.cover-art{width:104mm;height:104mm;object-fit:contain;align-self:center;margin:24px 0}.cover-identity{border-top:1px solid var(--line);padding-top:24px;margin-top:auto}.cover-name{font-size:40px;line-height:1.3;font-weight:700;color:var(--ink)}.cover-class{font-size:16px;margin-top:8px}.cover-date{font-size:14px;margin-top:24px;color:var(--muted)}
.matrix-axis{font-size:12px;color:var(--muted);margin-bottom:8px}.matrix{display:grid;gap:16px}.matrix-item{break-inside:avoid}.matrix-item h3{font-size:18px;margin-bottom:8px}.triad{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:16px}.triad dt{font-size:12px;font-weight:700;color:var(--muted);margin-bottom:4px}.triad dd{font-size:14px;line-height:1.6}.result{margin-top:8px;padding:8px 12px;background:var(--blue);border-radius:8px;display:grid;grid-template-columns:60px minmax(0,1fr);gap:12px;align-items:start}.result-label{font-size:12px;color:var(--muted);line-height:24px}.result-value{font-size:14px;line-height:24px;color:var(--ink);font-weight:700}.choice-list{display:flex;gap:16px;align-items:center;flex-wrap:wrap}.choice{display:inline-flex;align-items:center;gap:6px;color:var(--muted);font-weight:400;font-size:12px;line-height:24px}.choice .mark{display:inline-block;width:13px;height:13px;border:1px solid #617b93;border-radius:3px;line-height:12px;text-align:center;font-size:11px;color:#fff}.choice.selected{color:var(--ink);font-weight:700}.choice.selected .mark{background:var(--ink);border-color:var(--ink)}
.student-meta{display:flex;gap:32px;margin:0 0 24px;font-size:14px}.student-meta b{color:var(--ink)}.quote-grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:16px}.quote-card{background:var(--blue);border-radius:12px;padding:16px}.quote-card:nth-child(2){background:var(--violet)}.quote-card:nth-child(3){background:var(--green)}.quote-card h3{font-size:16px;margin-bottom:12px}.quote-card blockquote{font-size:14px;line-height:1.7;margin-top:8px}.quote-card .context{display:block;font-size:12px;color:var(--muted);margin-bottom:4px}.rows-head,.evidence-row{display:grid;grid-template-columns:160px minmax(0,1fr);gap:24px}.rows-head{font-size:12px;color:var(--muted);margin-bottom:8px}.evidence-row{padding:16px 0;border-top:1px solid var(--line);break-inside:avoid}.evidence-row h3{font-size:16px;line-height:1.6}.dimension-main{display:block;font-size:24px;line-height:1.2;margin-bottom:8px}.dimension-note{display:block;font-size:14px;white-space:nowrap;line-height:1.5}.evidence-row p{font-size:15px;line-height:1.75}.three-rows .evidence-row{padding:20px 0}.three-rows .evidence-row p{font-size:16px}.light-fields{display:grid;grid-template-columns:1fr 1fr;gap:24px 32px}.light-fields .field{padding:0}.light-fields dt{font-size:14px}.light-fields dd{font-size:16px}.synthesis{margin-top:24px;padding:20px 24px;background:var(--green);border-radius:12px}.synthesis dt{color:var(--green-ink);margin-bottom:8px}.synthesis dd{font-size:16px;color:var(--body)}.synthesis .template-note{color:var(--green-ink);margin-top:12px}.watch-help{font-size:14px;color:var(--muted);margin-bottom:24px}.watch-list{display:grid;gap:24px}.watch-item{display:grid;grid-template-columns:32px minmax(0,1fr);gap:16px}.watch-number{font-size:20px;font-weight:700;color:#315d86;line-height:1.4}.watch-item h3{font-size:16px;margin-bottom:8px}.watch-item p{font-size:16px;line-height:1.75}.message-box{padding:24px;background:var(--blue);border-radius:12px}.message-box .field+.field{margin-top:24px}.traits{display:flex;gap:12px;flex-wrap:wrap}.traits span{font-size:16px;font-weight:700;color:var(--ink)}.message-box dt{font-size:14px}.message-box dd{font-size:16px}.sentence-fill{font-size:16px;line-height:1.75}.sentence-fill strong{font-weight:700;color:var(--ink)}
@page{size:A4 portrait;margin:0}@media print{body{background:#fff}.toolbar{display:none}.page{margin:0;box-shadow:none;break-after:page;print-color-adjust:exact;-webkit-print-color-adjust:exact}.page:last-child{break-after:auto}}
@media screen and (max-width:830px){.toolbar{padding:8px 16px}.toolbar strong{font-size:12px}.toolbar button{padding:12px 16px;font-size:14px;white-space:nowrap}.page{width:calc(100% - 20px);height:auto;margin:10px;padding:32px 24px 56px}.page footer{right:24px;bottom:20px}.cover{min-height:780px}.cover h1{font-size:40px}.cover-art{width:100%;height:auto;max-width:340px;aspect-ratio:1}.cover-name{font-size:32px}.cover .school{font-size:16px}.triad,.quote-grid,.light-fields{grid-template-columns:1fr}.triad{gap:12px}.triad dd,.result-value,.quote-card blockquote,.evidence-row p{font-size:16px}.matrix{gap:32px}.matrix-item h3{font-size:20px;margin-bottom:16px}.result{grid-template-columns:1fr;gap:4px;margin-top:16px;padding:12px 16px}.choice-list{gap:8px 16px}.choice{font-size:14px}.student-meta{flex-direction:column;gap:8px}.evidence-row{grid-template-columns:1fr;gap:8px}.rows-head{display:none}.evidence-row h3{font-size:18px}.light-fields{gap:24px}.watch-item{gap:8px}.message-box{padding:20px}.template-note{font-size:12px}.section-heading{margin-bottom:24px}}
'''


def render_html(name,d):
    def page(n,body,cls=''):
        return f'<section class="page {cls}" id="page-{n}" aria-label="第{n}页"><div class="page-content">{body}</div><footer>{n:02d} / 05</footer></section>'
    _,uri=image_uri(d)
    # 封面仍沿用上一版配色与形象，只展示识别信息。
    cover=f'<div class="school">川大附小</div><h1>成像工具四：<span>AI信息汇聚</span></h1><p class="sources">家长问卷 × 学生访谈 × 教师观察</p><img class="cover-art" src="{uri}" width="1254" height="1254" alt="一年级学生主题虚拟形象"><div class="cover-identity"><p class="cover-name">{e(name)}</p><p class="cover-class">{e(d["class"])}</p><p class="cover-date">成像日期：{e(d["date"])}</p></div>'
    cover_page=f'<section class="page cover {d["gender"]}" id="page-1" aria-label="封面">{cover}<footer>01 / 05</footer></section>'
    cards=[]
    for i,m in enumerate(d['matrix']):
        triad=''.join(f'<div><dt>{e(label)}</dt><dd>{e(m[key])}</dd></div>' for key,label in zip(['parent','child','teacher'],MATRIX_LABELS[1:4]))
        if i==0:
            choices=''.join(f'<span class="choice {"selected" if option==m["insight"] else ""}"><span class="mark" aria-hidden="true">{"✓" if option==m["insight"] else ""}</span><span>{e(option)}</span></span>' for option in S['judgment_options'])
            value=f'<div class="choice-list" aria-label="初步判断：{e(m["insight"])}" data-selected="{e(m["insight"])}">{choices}</div>'
        else:value=f'<p class="result-value">{e(m["insight"])}</p>'
        cards.append(f'<article class="matrix-item"><h3>{e(m["dim"])}</h3><dl class="triad">{triad}</dl><div class="result"><span class="result-label">初步判断</span>{value}</div></article>')
    matrix=heading(H1)+f'<p class="matrix-axis">{e(MATRIX_LABELS[0])}</p><div class="matrix">'+''.join(cards)+'</div>'
    quotes=[]
    for key,label in zip(['parent','child','teacher'],SOURCE_LABELS):
        qt=''
        for q in d['quotes'][key]:
            if isinstance(q,dict):qt+=f'<blockquote><span class="context">{e(q["context"])}</span>“{e(q["quote"])}”</blockquote>'
            else:qt+=f'<blockquote>“{e(q)}”</blockquote>'
        quotes.append(f'<article class="quote-card"><h3>{e(label)}</h3>{qt}</article>')
    def evidence(rows,table_i,cls=''):
        heads=S['tables'][table_i][0]
        result=f'<div class="evidence-rows {cls}"><div class="rows-head"><span>{e(heads[0])}</span><span>{e(heads[1])}</span></div>'
        for label,value in rows:
            label_html=(f'<span class="dimension-main">{e(label[0])}</span><span class="dimension-note">{e(label[1:])}</span>' if table_i==1 else e(label))
            result+=f'<article class="evidence-row"><h3>{label_html}</h3><p>{e(value)}</p></article>'
        return result+'</div>'
    card=heading(H2)+f'<div class="student-meta"><p>学生姓名：<b>{e(name)}</b></p><p>成像日期：<b>{e(d["date"])}</b></p></div><div class="quote-grid">'+''.join(quotes)+'</div>'+heading(FIVE).replace('section-heading','section-heading sub')+evidence(d['five'],1)
    light=heading(THREE)+evidence(d['three'],2,'three-rows')+heading(LIGHT).replace('section-heading','section-heading sub')
    light+='<dl class="light-fields">'+''.join(field(label,d['light'][key]) for label,key in zip(LIGHT_LABELS,['want','strong','parent','teacher']))+'</dl>'
    light+=f'<dl class="synthesis">{field(LIGHT_RESULT,d["light"]["synthesis"])}<p class="template-note">{e(LIGHT_NOTE)}</p></dl>'
    watch=heading(WATCH)+f'<p class="watch-help">{e(WATCH_HELP)}</p><div class="watch-list">'
    for i,w in enumerate(d['watch'],1):watch+=f'<article class="watch-item"><span class="watch-number">{i:02d}</span><div><h3>{e(w["title"])}</h3><p>{e(w["text"])}</p></div></article>'
    watch+='</div>'+heading(MESSAGE).replace('section-heading','section-heading sub')
    traits=' '.join(c+t for c,t in zip('①②③',d['message']['traits']))
    watch+='<div class="message-box"><dl>'+field(MESSAGE_LABELS[0].split('①')[0],traits)+field(MESSAGE_LABELS[1].split('________')[0],d['message']['need'])+'</dl>'
    parts=MESSAGE_LABELS[2].split('________')
    watch+=f'<p class="sentence-fill" style="margin-top:24px">{e(parts[0])}<strong>{e(d["message"]["begin"])}</strong>{e(parts[1])}</p></div>'
    html='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><meta name="robots" content="noindex,nofollow"><title>'+e(name)+'｜'+e(TITLE)+'</title><style>'+CSS+'</style></head><body><nav class="toolbar" aria-label="报告操作"><strong>'+e(name)+' · 工具四</strong><button type="button" onclick="window.print()">导出 PDF / 打印</button></nav><main>'+cover_page+page(2,matrix)+page(3,card)+page(4,light)+page(5,watch)+'</main></body></html>'
    (OUT/f'{name}_工具四_初始成像.html').write_text(html)


def render_markdown(name,d):
    out=[f'# {TITLE}','',f'学生姓名：{name}　成像日期：{d["date"]}','',f'## {H1}','']
    for m in d['matrix']:
        out.extend([f'### {m["dim"]}',''])
        for label,key in zip(MATRIX_LABELS[1:4],['parent','child','teacher']):out.append(f'- {label}：{m[key]}')
        out.extend(['- 初步判断：'+m['insight'],''])
    out.extend([f'## {H2}','',f'学生姓名：{name}　成像日期：{d["date"]}',''])
    for key,label in zip(['parent','child','teacher'],SOURCE_LABELS):
        out.extend([f'### {label}','']+[quote_text(q) for q in d['quotes'][key]]+[''])
    for h,rows,table_i in [(FIVE,d['five'],1),(THREE,d['three'],2)]:
        out.extend([f'### {h}','',f'| {S["tables"][table_i][0][0]} | {S["tables"][table_i][0][1]} |','|---|---|'])
        out.extend('| '+label+' | '+value+' |' for label,value in rows);out.append('')
    out.extend([f'### {LIGHT}',''])
    for label,key in zip(LIGHT_LABELS,['want','strong','parent','teacher']):out.extend([label+d['light'][key],''])
    out.extend([LIGHT_RESULT,d['light']['synthesis'],'',LIGHT_NOTE,'',f'### {WATCH}','',WATCH_HELP,''])
    for w in d['watch']:out.extend([f'**{w["title"]}**：{w["text"]}',''])
    out.extend([f'### {MESSAGE}',''])
    for t in message_texts(d):out.extend([t,''])
    (OUT/f'{name}_工具四_文本.md').write_text('\n'.join(out))


def font(run,size=10.5,bold=False,color='263E55'):
    run.font.name='PingFang SC';run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
def para(doc,text='',size=10.5,bold=False,color='263E55',after=6):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(after);p.paragraph_format.line_spacing=1.35;font(p.add_run(text),size,bold,color);return p
def doc_heading(doc,text,size=18):
    i=text.find('（');main,note=(text[:i],text[i:]) if i>=0 else (text,'')
    p=para(doc,main,size,True,'12365C',5);p.paragraph_format.keep_with_next=True
    if note:para(doc,note,9,False,'50677C',10).paragraph_format.keep_with_next=True

def table(doc,heads,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(heads));t.autofit=False
    for c,label in zip(t.rows[0].cells,heads):
        font(c.paragraphs[0].add_run(label),9,True,'50677C')
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF5FC');c._tc.get_or_add_tcPr().append(sh)
    for row in rows:
        cells=t.add_row().cells
        for i,(c,value) in enumerate(zip(cells,row)):
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(8);p.paragraph_format.space_before=Pt(8);p.paragraph_format.line_spacing=1.4
            font(p.add_run(value),10.5,i==0,'12365C' if i==0 else '263E55')
    for row in t.rows:
        no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
        if widths:
            for c,w in zip(row.cells,widths):c.width=Mm(w)
    return t

def render_docx(name,d):
    doc=Document();sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(16);sec.bottom_margin=Mm(16);sec.left_margin=Mm(16);sec.right_margin=Mm(16)
    normal=doc.styles['Normal'];normal.font.name='PingFang SC';normal.font.size=Pt(10.5)
    para(doc,'川大附小',15,True,'50677C',28);para(doc,'成像工具四：',36,True,'12365C',0);para(doc,'AI信息汇聚',36,True,'12365C',12)
    para(doc,'家长问卷 × 学生访谈 × 教师观察',11,False,'50677C',12)
    art,_=image_uri(d);p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(art),width=Mm(98))
    para(doc,name,28,True,'12365C',6);para(doc,d['class'],12,False,'50677C',12);para(doc,'成像日期：'+d['date'],10,False,'50677C')
    doc.add_page_break();doc_heading(doc,H1)
    # Word 保留原表五列结构，内容与网页版使用同一字段。
    table(doc,MATRIX_LABELS,[[m['dim'],m['parent'],m['child'],m['teacher'],m['insight']] for m in d['matrix']],[27,37,37,37,40])
    doc.add_page_break();doc_heading(doc,H2)
    para(doc,f'学生姓名：{name}　成像日期：{d["date"]}',10.5,True,'12365C')
    table(doc,SOURCE_LABELS,[[ '\n'.join(quote_text(q) for q in d['quotes'][key]) for key in ['parent','child','teacher']]],[59,59,60])
    para(doc,'',after=6);doc_heading(doc,FIVE,15);table(doc,S['tables'][1][0],d['five'],[42,136])
    doc.add_page_break();doc_heading(doc,THREE);table(doc,S['tables'][2][0],d['three'],[42,136])
    para(doc,'',after=6);doc_heading(doc,LIGHT,15)
    for label,key in zip(LIGHT_LABELS,['want','strong','parent','teacher']):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(9);font(p.add_run(label),10.5,True,'12365C');font(p.add_run(d['light'][key]))
    para(doc,LIGHT_RESULT,11,True,'17644E',3);para(doc,d['light']['synthesis']);para(doc,LIGHT_NOTE,9,False,'50677C')
    doc.add_page_break();doc_heading(doc,WATCH);para(doc,WATCH_HELP,10,False,'50677C',16)
    for i,w in enumerate(d['watch'],1):
        para(doc,f'{i:02d}　{w["title"]}',12,True,'12365C',5);para(doc,w['text'],11,after=16)
    doc_heading(doc,MESSAGE,16)
    for t in message_texts(d):para(doc,t,11,after=16)
    doc.save(OUT/f'{name}_工具四_初始成像.docx')

if __name__=='__main__':
    validate()
    for name,d in DATA.items():
        render_html(name,d);render_docx(name,d);render_markdown(name,d)
        print(name,'原表字段核对通过；HTML / Word / 文本已更新')

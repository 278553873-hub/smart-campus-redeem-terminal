from pathlib import Path
import hashlib
import zipfile
import base64
import html
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin, section.bottom_margin = Cm(1.7), Cm(1.6)
section.left_margin, section.right_margin = Cm(1.9), Cm(1.9)
for name in ['Normal', 'Heading 1', 'Heading 2', 'Title']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
    style.font.color.rgb = RGBColor.from_string('171513')
doc.styles['Normal'].font.size = Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.15
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
header = section.header.paragraphs[0]
header.text = '教师手机端  /  班级页增量优化  /  开发交付'
header.style = doc.styles['Normal']
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string('6D6764')
footer = section.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('2026-10-23  ·  第 ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.add_run(' 页')

def p(text):
    return doc.add_paragraph(text)

def h(text):
    doc.add_heading(text, 2)

def page(title):
    if len(doc.paragraphs) > 0:
        doc.add_page_break()
    doc.add_heading(title, 1)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    if widths:
        for col, width in zip(t.columns, widths): col.width = Cm(width)
    for cell, title in zip(t.rows[0].cells, headers):
        cell.text = title
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F1EEEC'); cell._tc.get_or_add_tcPr().append(shade)
        for r in cell.paragraphs[0].runs: r.bold = True
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row): cell.text = value
    for row in t.rows:
        avoid = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(avoid)
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(4)
                para.paragraph_format.space_before = Pt(4)
                for run in para.runs: run.font.size = Pt(9.5)
    return t

doc.add_heading('班级页筛选与统计\n开发交付说明', 0)
p('年级筛选 · 任教班级复选 · 班级排行榜入口 · 数量统计')
p('版本：V1.0　｜　日期：2026-10-23　｜　基线：当前本地项目实现')
h('01  本次交付范围')
p('生产环境已按项目优化班级／社团与团队的页签切换、页面背景渐变和列表卡片。本次在这些成果上补齐筛选工具区、排行榜入口及统计展示规则，已有切换、渐变和卡片效果继续沿用。')
p('主要适用学校来源的教师手机端班级页；同时明确社团与团队页的数量统计及个人来源差异。排行榜内部排名算法、结算周期和榜单详情不在本次调整范围。')
h('页面结构与布局')
table(['层级', '内容与规则'], [
    ('标题栏', '沿用“班级 / 社团与团队”页签；筛选行紧接标题栏，不新增12像素间距。'),
    ('筛选工具行', '左侧：年级筛选＋任教班级；右侧：班级排行榜。最小高度44像素，两端对齐，不添加卡片、底色或阴影。'),
    ('数量统计行', '位于工具行下方、列表滚动区顶部。固定18像素高，展示“共N个”，不再占用44像素操作行。'),
    ('列表区域', '统计行与第一张卡片间距8像素。卡片、背景和既有页签动效继续沿用。'),
], [3.2,14])
p('横向对齐：学校班级页年级文字、社团页“我参与的”勾选框和两页“共N个”的左起点均为距屏幕左边24像素；卡片外边距仍为16像素，卡片正文起点为36像素。')
p('窄屏空间不足时允许工具行自然换行，排行榜保持右对齐；不得通过缩小字体、压缩触控区或遮挡微信胶囊来强行塞入一行。')
h('阅读口径')
p('本文尺寸均为逻辑显示像素，不是图片原始分辨率；微信小程序接入时按逻辑像素换算，避免将160px直接当成160rpx。图示仅用于结构说明，实际页面沿用已上线背景与卡片。')

page('02  年级筛选')
h('展示与交互')
p('仅在学校来源的“班级”页签展示。首次进入默认“全部年级”；点击打开公共底部弹窗，标题为“选择年级”。单选一个选项后立即关闭弹窗并刷新列表与数量，不增加确认或清空按钮。关闭弹窗而未选择时保持原值。')
table(['项目', '交付规格'], [
    ('触发器', '宽96、高44像素；文字与向下箭头左对齐紧邻，间距4像素。'),
    ('文字', '13像素、字重500；“全部年级”和具体年级统一主文字色 #171513。全部年级是有效范围值，不是灰色占位符。'),
    ('箭头', '16×16像素；三级文字色 #7B7572，不随选中年级变化。'),
    ('外观与反馈', '透明底，无常驻边框和阴影；触摸不增加底色或缩放。保留键盘焦点标记及弹窗展开状态的读屏反馈。'),
    ('弹窗选项', '全部年级独立在最上方；具体年级三列排列。选项可见高度44像素、触控高度48像素，胶囊形。'),
    ('选项状态', '当前项浅红底 #FFF1F1、红色边框和文字 #E02727；其他项白底、浅边框 #E7E2DF、次级文字 #6D6764。'),
], [3.2,14])
h('选项来源')
p('优先使用当前学校的学段配置；学校类型缺失时，按当前可访问班级的年级去重生成选项。不要用已经筛选后的结果反向生成选项，以免选项消失。')
table(['学校类型', '可选年级（均包含全部年级）'], [
    ('小学／初中／高中', '分别为一年级至六年级／七年级至九年级／高一至高三。'),
    ('九年一贯制', '小学、初中分组；一年级至九年级。'),
    ('十二年一贯制', '小学、初中、高中分组；一年级至九年级及高一至高三。'),
    ('完全中学', '初中、高中分组；七年级至九年级及高一至高三。'),
], [3.2,14])
p('多学段学校展示对应学段名称；单学段不重复展示学段标题。选定年级只改变班级范围，不改变“任教班级”的勾选状态。')

page('03  任教班级复选与筛选联动')
h('展示与交互')
p('仅在学校来源的“班级”页签展示，默认不勾选。勾选框和“任教班级”文字组成完整点击区域，点任意位置切换；立即更新列表和统计，无需提交。')
table(['项目', '交付规格'], [
    ('触控区域', '高度44像素，左右内边距8像素；勾选框与文字间距4像素。'),
    ('勾选框', '16×16像素、圆角5像素；未选中白底、1像素边框 #918985。'),
    ('选中状态', '仅框体变为品牌红 #E02727，内部白色勾号12×12像素；不将整项填充为红色。'),
    ('文字', '13像素、字重500、主文字色 #171513，选中与未选中一致。'),
    ('交互反馈', '整项不添加按压底色、阴影或缩放；通过勾选状态及结果变化反馈。读屏可识别选中状态，键盘可操作。'),
], [3.2,14])
h('“任教”的数据口径')
p('任教班级集合＝学科任教班级 ∪ 班主任班级 ∪ 副班主任班级，按班级唯一标识去重。一个班级即使同时有多门学科和多种任职关系，也只展示一次、统计一次。学校管理者身份本身不等于任教关系。')
table(['年级条件', '任教复选', '最终展示范围'], [
    ('全部年级', '未选中', '当前来源下可访问的全部班级'),
    ('具体年级', '未选中', '可访问班级中属于该年级的班级'),
    ('全部年级', '已选中', '可访问班级中具有上述任教关系的班级'),
    ('具体年级', '已选中', '同时属于该年级、且具有上述任教关系的班级'),
], [3.2,3,11])
h('状态与位置')
p('年级与任教条件取交集。取消任教筛选保留年级；改回全部年级保留任教勾选。条件变化后，列表滚动位置回到首项，顶部导航位置保持不动。')
p('同一班级列表实例内切换“班级 / 社团与团队”保留各自条件，班级条件不影响团队列表。切换来源空间时，年级恢复全部、任教及我参与的取消勾选、筛选弹窗关闭。当前实现未持久化筛选条件，整页刷新或组件重新挂载后恢复默认；本次不新增跨会话记忆。')

page('04  班级排行榜入口与图标交付')
p('使用附件中的正式透明图片。图内已经包含奖杯和“班级排行榜”文字，页面不再重复叠加文字、箭头、边框、底板或阴影。')
doc.add_picture(str(ROOT / '素材/班级排行榜.png'), width=Cm(13.6))
p('正式素材预览（上图放大用于核对；页面实际使用下列尺寸）')
table(['项目', '交付规格'], [
    ('展示条件', '学校来源＋当前为班级页签＋学校开启班级记录能力（classRecordEnabled 为 true）。团队页签、个人来源和未开启能力时不展示。'),
    ('放置位置', '筛选工具行右侧，与年级、任教控件同一行；不移到标题栏、我的页面或悬浮按钮。窄屏换行后仍右对齐。'),
    ('可点击区域', '130×44像素；整块区域可点击，不只是图片有色区域。'),
    ('图片容纳区域', '124×24像素，居中、等比完整显示。不可裁切或强行拉伸为124×24。'),
    ('实际等比尺寸', '素材比例约5.80:1，放入124×24框后实际绘制约124×21.4像素，上下自然留白。'),
    ('点击行为', '进入既有班级排行榜页面；当前入口不向榜单传递年级或任教条件，不自动把榜单变成筛选结果排名。'),
    ('无障碍', '入口读屏名称“查看班级排行榜”；图片作为装饰避免重复朗读；保留键盘焦点状态。'),
], [3.2,14])
h('交付资源')
p('文件：素材/班级排行榜.png\n格式：透明 PNG　｜　原始尺寸：2128×367像素\n来源：mobile-app/assets/resources/class-leaderboard-entry-sunset-trophy-art.png\n项目映射：ASSETS.CLASS.LEADERBOARD_ENTRY')
p('直接使用原图或由正常资源流水线生成等比压缩版本，不重新绘制、抠图、添加白底或使用其他奖杯备选图。图片为空或筛选结果为0，不应据此隐藏仍具有访问权限的排行榜入口。')

page('05  统计数据与空结果展示')
h('数量统计口径')
table(['场景', '文案与计算'], [
    ('学校来源 · 班级', '共N个；N为年级和任教条件共同作用后的班级数量。每个班级计1个，不是学生数、任教关系条数或未筛选的全校总数。'),
    ('社团与团队', '共N个；N为当前可见团队数量。启用“我参与的”后，仅计算当前教师负责或协作的团队，按团队去重。'),
    ('个人来源 · 班级', '无隐藏班级时“共N个班级”；存在隐藏项时“显示V/T个班级”，V为可见数量，T为该来源班级总数。保留既有班级管理行，本次不套用学校统计行。'),
], [4,13.2])
p('学校班级页与团队页最终文案统一为“共N个”，例如“共12个”，不加空格、不再附加“班级”或“社团与团队”。旧规范中“共x个班级”的学校口径已落后于当前代码及验收断言，本次交付以本文为准；个人来源继续保留“班级”后缀。')
h('学校班级与团队页样式')
p('12像素、字重500、等宽数字、次级文字色 #6D6764；行高与占位高度均为18像素。左起点24像素，与筛选起点对齐；与首张卡片间距8像素。不要改为徽标、白色卡片或独立44像素行。筛选完成后数量与列表同步更新，读屏温和播报变化。')
h('加载与异常')
table(['状态', '展示规则'], [
    ('首次请求，无成功结果', '数量原位置显示48×10像素骨架占位，与卡片共用加载效果；不能先显示“共0个”。'),
    ('已成功加载／切换筛选', '展示真实数量；筛选不重播首次加载动效，不叠加人为等待。'),
    ('后台刷新，有成功缓存', '保留缓存的列表和数量；成功后同次更新。已确认的空结果也属于成功缓存。'),
    ('成功结果为空', '显示“共0个”和缺省态。学校班级文案“没有符合条件的班级”，缺省图显示框统一160×160像素。'),
    ('失败', '首次失败进入现有错误反馈与重试，不伪装成零结果；刷新失败保留成功缓存。生产数据层需完成错误分支接入。'),
], [4,13.2])
p('若生产环境采用分页，N应为当前筛选条件下的总数，不是已加载页的条数。当前演示对完整列表计算长度；接口接入时须返回同口径总数，并避免旧请求覆盖新筛选结果。')

page('06  验收清单与实现索引')
table(['验收项', '预期结果'], [
    ('默认进入', '全部年级、任教未选中；数量等于当前学校可访问班级总数。'),
    ('选择具体年级', '单选即生效并关窗；文字保持黑色；数量与结果一致。'),
    ('任教交集', '选定年级再勾任教，仅剩交集；取消勾选不清空年级。'),
    ('任职去重', '同班兼任班主任、副班主任或多学科任教，只计1个。'),
    ('零结果', '共0个＋160像素缺省图；保留年级、任教操作，允许恢复范围。'),
    ('排行榜权限', '学校开启班级记录且处于班级页签时显示，关闭能力或切到团队时隐藏。'),
    ('排行榜素材', '完整、透明、无拉伸；点击区域130×44，图框124×24。'),
    ('页签与来源', '班级筛选不污染团队；同实例切页签保留条件，切来源重置。'),
    ('团队统计', '“我参与的”启用后，只统计本人负责或协作的团队。'),
    ('加载与失败', '首载无假零值；缓存刷新不闪空；请求失败不当作空结果。'),
    ('小屏与读屏', '320／375／393／430像素宽检查不遮挡、不缩小热区；可读出筛选与复选状态。'),
    ('既有效果回归', '已交付的页签切换、背景渐变、卡片和加载效果保持一致。'),
], [3.2,14])
h('代码定位（均相对项目根目录）')
p('页面布局、筛选、统计：mobile-app/views/ClassListView.tsx\n年级弹窗：mobile-app/components/ui/MobileGradePickerSheet.tsx\n学校学段、排行榜能力：mobile-app/domain/teacherSpaceAccess.ts\n尺寸与颜色：mobile-app/styles/teacherMobileTokens.ts\n空状态：mobile-app/components/ui/MobileEmptyState.tsx\n图片映射：mobile-app/assets/images.ts')
h('接入边界')
p('前端筛选不能代替服务端权限控制；输入班级与团队集合须先限定为当前账号、当前来源可访问的数据。本交付定义展示和筛选口径，不要求重建已上线榜单或权限体系。生产首次失败、分页总数和请求竞态需结合现有接口实现。')
p('加载细节参见 docs/演示项目/教师手机端/班级列表加载动效开发交付说明.md。本文按2026-10-23本地代码核对，未对生产环境进行在线验收。')

output = ROOT / '教师手机端-班级页筛选与统计-开发交付说明.docx'
doc.save(output)
from docx.text.paragraph import Paragraph
from docx.table import Table
blocks = []
for element in doc.element.body:
    if element.tag == qn('w:p'):
        paragraph = Paragraph(element, doc)
        if element.xpath('.//w:br[@w:type="page"]'):
            blocks.append('<hr>')
        if element.xpath('.//w:drawing'):
            encoded = base64.b64encode((ROOT / '素材/班级排行榜.png').read_bytes()).decode()
            blocks.append(f'<figure><img alt="班级排行榜正式图标" src="data:image/png;base64,{encoded}"></figure>')
        if paragraph.text:
            tag = 'h1' if paragraph.style.name in ['Title', 'Heading 1'] else 'h2' if paragraph.style.name == 'Heading 2' else 'p'
            blocks.append(f'<{tag}>' + html.escape(paragraph.text).replace('\n', '<br>') + f'</{tag}>')
    elif element.tag == qn('w:tbl'):
        rows = Table(element, doc).rows
        blocks.append('<div class="table-wrap"><table>')
        for i,row in enumerate(rows):
            tag = 'th' if i == 0 else 'td'
            blocks.append('<tr>' + ''.join(f'<{tag}>{html.escape(c.text)}</{tag}>' for c in row.cells) + '</tr>')
        blocks.append('</table></div>')
preview = ROOT / '教师手机端-班级页筛选与统计-开发交付说明.html'
preview.write_text('''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>班级页筛选与统计 · 开发交付</title><style>
*{box-sizing:border-box}body{margin:0;background:#f8f6f5;color:#171513;font:15px/1.8 -apple-system,BlinkMacSystemFont,"PingFang SC","Microsoft YaHei",sans-serif}main{max-width:960px;margin:32px auto;padding:40px 48px;background:white}h1{font-size:28px;line-height:1.4;margin:0 0 20px}h2{font-size:18px;margin:28px 0 10px}p{margin:12px 0;overflow-wrap:anywhere}hr{border:0;border-top:1px solid #e7e2df;margin:48px 0}table{border-collapse:collapse;width:100%;font-size:14px}td,th{padding:12px 14px;border:1px solid #e7e2df;text-align:left;vertical-align:top}th{background:#f1eeec}td:first-child{min-width:120px;width:24%}.table-wrap{overflow-x:auto}figure{margin:24px 0}figure img{width:100%;max-width:650px;height:auto}@media(max-width:640px){main{margin:0;padding:24px 18px}h1{font-size:24px}td,th{padding:10px}}@media print{body{background:white}main{margin:0;padding:0}hr{break-before:page;border:0}tr{break-inside:avoid}h1,h2{break-after:avoid}}</style><main>''' + ''.join(blocks) + '</main></html>')
readme = '交付内容\n\n1. Word文档：功能规则、尺寸、统计口径和验收清单，文内已嵌入正式排行榜图标。\n2. 网页预览：双击HTML文件即可阅读，图片已内嵌，可离线打开。\n3. 素材/班级排行榜.png：2128×367透明原图，可直接交开发。\n\n范围：教师手机端班级页增量优化，沿用已完成的页签切换、背景渐变及卡片效果。\n'
(ROOT / '阅读说明.txt').write_text(readme)
with zipfile.ZipFile(ROOT / '班级页筛选与统计-开发交付包.zip', 'w', zipfile.ZIP_DEFLATED) as z:
    for f in [output, preview, ROOT / '阅读说明.txt', ROOT / '素材/班级排行榜.png']:
        z.write(f, f.relative_to(ROOT))
print(output)
print('图标 SHA256:', hashlib.sha256((ROOT / '素材/班级排行榜.png').read_bytes()).hexdigest())
